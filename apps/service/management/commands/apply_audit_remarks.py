import re
from collections import Counter, defaultdict

import openpyxl
from django.core.management.base import CommandError
from django.db import transaction
from django.db.models import Max

from apps.service.models import Product, ProductItemCategory, ProductSubCategory

from .apply_audit_corrections import (
    COLUMN_ID, NAME_FIELDS, Command as CorrectionsCommand, clean, named, strip_articul,
)

# The audit workbook lists every remark on the sheet "Все замечания": which product, which field, what is wrong
# and what to do. The texts are generated from a few templates, so the remarks that say exactly what to do are
# executed here. A remark is skipped when the content team already corrected that cell by hand (yellow / green
# fill) or when the product no longer looks like it did at the audit.

REMARKS_SHEET = "Все замечания"
CREATE_MARK = "(создать)"

LATIN_TO_CYRILLIC = dict(zip("ABEKMHOPCTXaeopcyx", "АВЕКМНОРСТХаеорсух"))
CYRILLIC_TO_LATIN = {cyrillic: latin for latin, cyrillic in LATIN_TO_CYRILLIC.items()}

# names left in Turkish / transliterated Russian: only the wordings we are sure about
TRANSLIT = {
    "VIKO PILOT 6X1,5MT": "Viko Сетевой фильтр 6 гнёзд 1,5 м",
    "PILOT (SETEVOY FILTR)": "Сетевой фильтр",
    "Avtomat": "Автомат",
}


def is_cyrillic(char):
    return "а" <= char.lower() <= "я" or char in "ёЁ"


def quoted(text):
    match = re.search(r"«(.*?)»", text)
    return match.group(1) if match else ""


def keep_case(source, replacement):
    if source.isupper() and len(source) > 1:
        return replacement.upper()
    if source[:1].isupper():
        return replacement[:1].upper() + replacement[1:]
    return replacement


NEEDS_A_PERSON = object()


# ---- name remarks: (current name, remark) -> fixed name; None when the remark doesn't fit the name any more
# (fixed already or edited since the audit); NEEDS_A_PERSON when there is no sure wording ----

def fix_example(name, remark):
    # "Название начинается с бренда/серии" -> the ready name is given: Пример: «…»
    if clean(remark["Наименование"]) != name:
        return None
    example = clean(quoted(remark["Что сделать"]))
    # the example writes sizes with "×", the catalog (and the content team) keep the original letter
    separators = iter(re.findall(r"(?<=\d)[xXхХ](?=\d)", name))
    return re.sub("×", lambda match: next(separators, "х"), example)


def fix_typo(name, remark):
    # «ричаг» -> "Исправить → рычаг";  «Изоляцией на наконечнику» -> "Исправить → «Изоляция для наконечника …»"
    wrong = quoted(remark["Проблема"])
    right = quoted(remark["Что сделать"]) or remark["Что сделать"].split("→")[-1]
    right = clean(right.replace("…", ""))
    if not wrong or not right or wrong.lower() not in name.lower():
        return None
    return re.sub(re.escape(wrong), lambda match: keep_case(match.group(), right), name, flags=re.IGNORECASE)


def fix_mixed_alphabets(name, remark):
    # "Смешаны латиница и кириллица: 1ПCТ-10-35/50": look-alike letters go to the alphabet of the rest of the
    # word; when a letter has no twin ("розеткаLEE"), the two parts are simply separated
    token = clean(remark["Проблема"].split(":", 1)[-1])
    if not token or token not in name:
        return None
    cyrillic = [char for char in token if char.isalpha() and is_cyrillic(char)]
    latin = [char for char in token if char.isalpha() and not is_cyrillic(char)]
    minority, twins = (cyrillic, CYRILLIC_TO_LATIN) if len(cyrillic) < len(latin) else (latin, LATIN_TO_CYRILLIC)
    if all(char in twins for char in minority):
        fixed = "".join(twins.get(char, char) if char in minority else char for char in token)
    else:
        fixed = split_alphabets(token)
    return name.replace(token, fixed)


def split_alphabets(text):
    return re.sub(r"(?<=[а-яё])(?=[A-Za-z])|(?<=[A-Za-z])(?=[а-яё])", " ", text)


def fix_glued_words(name, remark):
    fixed = re.sub(r"(?<=[а-яё])(?=[A-Z])", " ", name)
    return fixed if fixed != name else None


def fix_caps(name, remark):
    # "УЗИП", "НШВИ" are abbreviations; a name is shouting only when it has a long word in capitals
    if not any(len(word) >= 5 for word in re.findall(r"(?<!\w)[А-ЯЁ]+(?!\w)", name)):
        return name
    fixed = re.sub(r"(?<!\w)[А-ЯЁ]+(?!\w)", lambda match: match.group().lower(), name)
    return fixed[:1].upper() + fixed[1:]


def fix_bracket(name, remark):
    return f"{name})" if name.count("(") == name.count(")") + 1 else None


def fix_kelvin(name, remark):
    # "Цветовая температура «3500» без K"
    value = quoted(remark["Проблема"])
    fixed = re.sub(rf"(?<!\d){re.escape(value)}(?![\dKК])", f"{value}K", name) if value else name
    return fixed if fixed != name else None


def fix_translit(name, remark):
    for wrong, right in TRANSLIT.items():
        if wrong in name:
            return name.replace(wrong, right)
        if right in name:
            return name
    return NEEDS_A_PERSON


def fix_spaces(name, remark):
    return name  # the name is already read without tabs / line breaks / double spaces


NAME_FIXES = (
    # (label, test on "Проблема" / "Что сделать", function); the first match wins
    ("name: brand/series moved to the end", lambda problem, action: action.startswith("Пример: «"), fix_example),
    ("name: typo", lambda problem, action: problem.startswith(("Опечатка/ошибка:", "Неграмотное название")), fix_typo),
    ("name: mixed alphabets", lambda problem, action: problem.startswith("Смешаны латиница и кириллица:"),
     fix_mixed_alphabets),
    ("name: glued words", lambda problem, action: problem.startswith("Слова слиплись"), fix_glued_words),
    ("name: capitals", lambda problem, action: "КАПСОМ" in problem, fix_caps),
    ("name: bracket", lambda problem, action: "непарные скобки" in problem, fix_bracket),
    ("name: colour temperature", lambda problem, action: problem.startswith("Цветовая температура"), fix_kelvin),
    ("name: transliteration", lambda problem, action: "Нет русского названия" in problem
     or "Название транслитом" in problem, fix_translit),
    ("name: tabs / line breaks", lambda problem, action: "табуляция" in problem, fix_spaces),
)


class Command(CorrectionsCommand):
    help = (
        'Execute the audit remarks of one sheet (the "Все замечания" sheet of the audit workbook) that say '
        "exactly what to do: articul out of the name, item category, brand, name format and typos. "
        "Cells the content team corrected by hand are left to apply_audit_corrections. Run with --dry-run first."
    )

    def handle(self, *args, **options):
        self.setup(options)
        self.sheet_name = options["sheet"]
        self.not_automated = Counter()

        remarks = self.read_remarks(options["file"])
        self.stdout.write(f"Open remarks of the sheet: {sum(len(items) for items in remarks.values())} "
                          f"on {len(remarks)} products")

        with transaction.atomic():
            products = {
                product["id"]: product for product in Product.objects.rewrite(False).filter(pk__in=remarks).values(
                    "id", *NAME_FIELDS, "articul_code", "brand_id", "brand__name_ru", "product_model__brand_id",
                    "product_item_category_id", "product_item_category__name_ru",
                    "product_item_category__product_sub_category__product_category_id",
                )
            }
            self.plan_articul_removal(remarks, products)
            for product_id, items in remarks.items():
                if product_id in products:
                    self.apply_remarks(products[product_id], items)
                else:
                    self.stats["products not found"] += 1
            if self.dry_run:
                transaction.set_rollback(True)

        self.finish()
        if self.not_automated:
            self.stdout.write("\nLeft for the content team (no exact instruction, or data is needed):")
            for (kind, action), count in self.not_automated.most_common():
                self.stdout.write(f"  {count:>5}  {kind}: {action}")

    def read_remarks(self, path):
        workbook = openpyxl.load_workbook(path, read_only=True)
        for sheet_name in (self.sheet_name, REMARKS_SHEET):
            if sheet_name not in workbook.sheetnames:
                raise CommandError(f"No sheet {sheet_name!r}. Sheets: {', '.join(workbook.sheetnames)}")

        rows = workbook[self.sheet_name].iter_rows()
        headers = [clean(cell.value) for cell in next(rows)]
        corrected = set()
        for row in rows:
            product_id = row[headers.index(COLUMN_ID)].value
            corrected.update((product_id, header) for header, cell in zip(headers, row) if self.is_corrected(cell))

        rows = workbook[REMARKS_SHEET].iter_rows(values_only=True)
        headers = [clean(value) for value in next(rows)]
        remarks = defaultdict(list)
        for row in rows:
            remark = {header: clean(value) for header, value in zip(headers, row)}
            if remark["Лист"] != self.sheet_name or not remark[COLUMN_ID].isdigit():
                continue
            product_id = int(remark[COLUMN_ID])
            if (product_id, remark["Поле"]) in corrected:
                self.stats["remarks on cells corrected by hand (skipped)"] += 1
            else:
                remarks[product_id].append(remark)
        workbook.close()
        return remarks

    def plan_articul_removal(self, remarks, products):
        # an articul leaves the name only if the name stays unique: several Schneider products differ by it alone
        candidates = {}
        for product_id, items in remarks.items():
            product = products.get(product_id)
            for remark in items:
                if product and remark["Тип ошибки"] == "Артикул":
                    stripped = strip_articul(clean(product["name_ru"] or product["name"]), quoted(remark["Проблема"]))
                    if stripped:
                        candidates[product_id] = stripped
        taken = Counter(candidates.values())
        taken.update(
            clean(name) for name in Product.objects.rewrite(False).exclude(pk__in=candidates).values_list(
                "name_ru", flat=True)
        )
        self.removable_articul = {product_id for product_id, name in candidates.items() if taken[name] == 1}

    def apply_remarks(self, product, remarks):
        current = clean(product["name_ru"] or product["name"])
        name = current
        changes = {}

        # a remark with a ready name replaces the whole name, so it goes before the ones that touch a word
        for remark in sorted(remarks, key=lambda remark: not remark["Что сделать"].startswith("Пример: «")):
            kind, problem, action = remark["Тип ошибки"], remark["Проблема"], remark["Что сделать"]

            if kind == "Артикул" and action.startswith("Перенести артикул"):
                name = self.move_articul(product, name, quoted(problem), changes)
            elif kind == "Категория" and action.startswith("Перенести в «"):
                self.move_to_category(product, remark, changes)
            elif kind == "Бренд" and action.startswith("Поставить:"):
                self.set_brand(product, remark, changes)
            elif kind == "Формат" and "начинается с кода" in problem:
                pass  # the code is the articul: done by the "Артикул" remark of the same product
            elif remark["Поле"] == "Наименование" and any(test(problem, action) for _, test, _ in NAME_FIXES):
                label, fix = next((label, fix) for label, test, fix in NAME_FIXES if test(problem, action))
                fixed = fix(name, remark)
                if fixed is NEEDS_A_PERSON:
                    self.count_not_automated(remark)
                elif fixed is None:
                    self.stats[f"{label}: fixed already / the name has changed since the audit"] += 1
                elif fixed == name:
                    self.stats[f"{label}: nothing to do"] += 1
                else:
                    self.stats[label] += 1
                    name = fixed
            else:
                self.count_not_automated(remark)

        if name != current:
            # the columns that repeat the Russian name follow it, a real translation stays
            changes.update({field: name for field in NAME_FIELDS if field == "name_ru" or product[field] == current})
            self.log(product, "name", current, name)
        if changes:
            Product.objects.rewrite(False).filter(pk=product["id"]).update(**changes)
            self.stats["products changed"] += 1

    def count_not_automated(self, remark):
        self.not_automated[(remark["Тип ошибки"], re.sub(r"«[^»]*»", "«…»", remark["Что сделать"]))] += 1

    def move_articul(self, product, name, code, changes):
        articul = clean(product["articul_code"])
        stripped = strip_articul(name, code)
        if not code or articul not in ("", code):
            self.stats["articul: the product already has another one (skipped)"] += 1
        elif stripped is None:
            # "(225X311X130)", "(RI-TECH)" in the middle of a name are sizes and series, not articuls
            self.stats["articul: already moved" if articul else "articul: not a separate code in the name (skipped)"] += 1
        else:
            if not articul:
                changes["articul_code"] = code
                self.log(product, "articul", "", code)
            if product["id"] in self.removable_articul:
                self.stats["articul: moved out of the name"] += 1
                return stripped
            self.stats["articul: filled, kept in the name (the only difference between products)"] += 1
        return name

    def set_brand(self, product, remark, changes):
        if clean(product["brand__name_ru"]) != remark["Текущее значение"]:
            self.stats["brand: changed since the audit (skipped)"] += 1
            return
        brand = self.get_brand(clean(remark["Что сделать"].split(":", 1)[1].split("(")[0]))
        if not brand or brand.pk == product["brand_id"]:
            self.stats["brand: nothing to do"] += 1
            return
        changes.update(self.brand_change(product, brand))
        self.stats["brand: changed"] += 1
        self.log(product, "brand", product["brand__name_ru"], brand.name)

    def move_to_category(self, product, remark, changes):
        if clean(product["product_item_category__name_ru"]) != remark["Текущее значение"]:
            self.stats["category: changed since the audit (skipped)"] += 1
            return
        item_category = self.item_category_by_path(product, quoted(remark["Что сделать"]))
        if not item_category:
            self.stats["category: target not found / ambiguous (skipped)"] += 1
            self.warn(f"product {product['id']}: no single item category for {remark['Что сделать']!r}")
        elif item_category.pk == product["product_item_category_id"]:
            self.stats["category: nothing to do"] += 1
        else:
            changes.update(self.item_category_change(product, item_category))
            self.stats["category: moved"] += 1
            self.log(product, "item category", product["product_item_category__name_ru"], item_category.name)

    def item_category_by_path(self, product, path):
        # "Электрика → Низковольтное оборудование → Дифференциальные автоматы (АВДТ) (создать)" or just the name
        parts = [clean(part) for part in path.replace(CREATE_MARK, "").split("→")]
        name, sub_category_name = parts[-1], parts[-2] if len(parts) > 1 else ""
        category_id = product["product_item_category__product_sub_category__product_category_id"]
        key = (name.lower(), sub_category_name.lower(), category_id)
        if key in self.item_categories and self.item_categories[key]:
            return self.item_categories[key]

        item_category = self.get_item_category(name, sub_category_name, category_id)
        if not item_category and CREATE_MARK in path and not named(ProductItemCategory, name).exists():
            sub_categories = list(named(ProductSubCategory, sub_category_name).filter(product_category_id=category_id))
            if len(sub_categories) == 1:
                position = ProductItemCategory.objects.filter(product_sub_category=sub_categories[0]).aggregate(
                    last=Max("position"))["last"] or 0
                # save() (not create()): CategoryMixin builds the slug there
                item_category = ProductItemCategory(
                    product_sub_category=sub_categories[0], name=name, name_ru=name, position=position + 1)
                item_category.save()
                self.item_categories[key] = item_category
                self.stats["item categories created"] += 1
                self.stdout.write(f"  Item category {item_category.pk} {name!r} created in {sub_category_name!r}")
        return item_category
