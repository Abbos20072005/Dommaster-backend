import re
from collections import Counter

import openpyxl
from django.core.cache import cache
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.db.models import Q

from apps.service.models import (
    Brand, Product, ProductAttributeValue, ProductFilterNumericValue, ProductItemCategory,
)

# The content team corrects a sheet of the audit workbook (an export of export_products_by_category with audit
# notes) and marks every cell it changed with a fill: yellow in the first round, green in the second. Only those
# cells are applied - the rest of the sheet is an old snapshot of the DB and must not overwrite it.
# A red fill is the team's own "problem" mark, not a correction.

NAME_FIELDS = ("name", "name_uz", "name_ru", "name_en")
EDITED_FILLS = {"FFFFFF00", "FF92D050"}

COLUMN_ID = "ID"
COLUMN_NAME = "Наименование"
COLUMN_SUB_CATEGORY = "Подкатегория"
COLUMN_ITEM_CATEGORY = "Предметная категория"
COLUMN_BRAND = "Бренд"
APPLIED_COLUMNS = (COLUMN_NAME, COLUMN_ITEM_CATEGORY, COLUMN_BRAND)


def clean(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def named(model, name):
    condition = Q()
    for field in NAME_FIELDS:
        condition |= Q(**{f"{field}__iexact": name})
    return model.objects.rewrite(False).filter(condition).order_by("id")


def strip_articul(name, code):
    """The name without the articul, or None when the code is not a separate piece at the start / end of it."""
    escaped = re.escape(code)
    for pattern in (rf"\s*\({escaped}\)$", rf"\s+{escaped}$", rf"^{escaped}\s+"):
        if re.search(pattern, name):
            return clean(re.sub(pattern, "", name)).strip(" ,")
    return None


def without_articul(name, articul):
    # the team types over an old export, so a corrected name may still end with the code that has moved to
    # articul_code since then; a mistyped code at that place is no better
    stripped = strip_articul(name, articul)
    return stripped if stripped is not None else clean(re.sub(r"[\s,]+\d{6,}$", "", name))


class Command(BaseCommand):
    help = (
        "Apply the cells the content team corrected (yellow / green fill) on one sheet of the audit workbook: "
        "product name, item category, brand. Run with --dry-run first."
    )

    def add_arguments(self, parser):
        parser.add_argument("file", help="Path to the audit workbook (.xlsx).")
        parser.add_argument("--sheet", required=True, help='Sheet to apply, e.g. "Электрика".')
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Do everything in a transaction and roll it back, nothing is saved.",
        )
        parser.add_argument(
            "--verbose",
            action="store_true",
            help="Print every change.",
        )

    def handle(self, *args, **options):
        self.setup(options)

        edits, other_columns = self.read_sheet(options["file"], options["sheet"])
        self.stdout.write(f"Rows with corrected cells: {len(edits)}")
        for column, count in other_columns.items():
            self.warn(f"{count} corrected cells in the column {column!r} are not applied (not supported)")

        with transaction.atomic():
            products = {
                product["id"]: product for product in Product.objects.rewrite(False).filter(pk__in=edits).values(
                    "id", *NAME_FIELDS, "articul_code", "brand_id", "product_model__brand_id",
                    "product_item_category_id", "product_item_category__product_sub_category__product_category_id",
                )
            }
            self.find_new_duplicates(edits, products)
            for product_id, cells in edits.items():
                product = products.get(product_id)
                if not product:
                    self.stats["products not found"] += 1
                    self.warn(f"product {product_id}: not found")
                    continue
                self.apply(product, cells)
            if self.dry_run:
                transaction.set_rollback(True)

        self.finish()

    def setup(self, options):
        self.dry_run = options["dry_run"]
        self.verbose = options["verbose"]
        self.stats = Counter()
        self.brands = {}
        self.item_categories = {}
        self.touched_item_categories = set()
        self.moved_to = set()
        if self.dry_run:
            self.stdout.write(self.style.WARNING("DRY RUN - everything below is rolled back at the end"))

    def finish(self):
        if not self.dry_run:
            # update() sends no signals; the brand list of an item category is cached too
            for item_category_id in self.touched_item_categories:
                cache.delete(f"product:available_filters:cat:{item_category_id}")
                cache.delete(f"categories:list:item:category={item_category_id}")

        self.stdout.write("")
        for label, count in self.stats.items():
            self.stdout.write(f"  {label}: {count}")
        if self.moved_to:
            ids = " ".join(str(pk) for pk in sorted(self.moved_to))
            self.warn(f"filters of the moved products: manage.py backfill_filter_data --item-category {ids}")
        self.stdout.write(self.style.SUCCESS("DRY RUN finished, nothing saved" if self.dry_run else "Done"))

    def warn(self, message):
        self.stdout.write(self.style.WARNING(f"  ! {message}"))

    @staticmethod
    def is_corrected(cell):
        return bool(cell.fill and cell.fill.fill_type == "solid" and cell.fill.fgColor.rgb in EDITED_FILLS)

    @staticmethod
    def corrected_name(product, cells):
        name = clean(cells[COLUMN_NAME])
        articul = clean(product["articul_code"])
        return without_articul(name, articul) if name and articul else name

    def find_new_duplicates(self, edits, products):
        # a corrected name must not make two different products look the same (a copy-paste slip in the sheet);
        # products that already share a name may keep sharing it
        current = {pk: clean(name) for pk, name in Product.objects.rewrite(False).values_list("id", "name_ru")}
        planned = dict(current)
        for product_id, cells in edits.items():
            if COLUMN_NAME in cells and product_id in products:
                planned[product_id] = self.corrected_name(products[product_id], cells) or current[product_id]
        owners = {}
        for product_id, name in planned.items():
            owners.setdefault(name, []).append(product_id)
        self.new_duplicates = {
            name for name, ids in owners.items() if len(ids) > 1 and len({current[pk] for pk in ids}) > 1
        }

    def read_sheet(self, path, sheet_name):
        # read_only: the workbook carries thousands of embedded pictures and cell notes we don't need
        workbook = openpyxl.load_workbook(path, read_only=True)
        if sheet_name not in workbook.sheetnames:
            raise CommandError(f"No sheet {sheet_name!r}. Sheets: {', '.join(workbook.sheetnames)}")
        rows = workbook[sheet_name].iter_rows()
        headers = [clean(cell.value) for cell in next(rows)]
        missing = [column for column in (COLUMN_ID, *APPLIED_COLUMNS) if column not in headers]
        if missing:
            raise CommandError(f"Columns not found on the sheet: {', '.join(missing)}")

        edits = {}
        other_columns = Counter()
        for row in rows:
            values = {header: cell.value for header, cell in zip(headers, row)}
            edited = [header for header, cell in zip(headers, row) if self.is_corrected(cell)]
            if not edited or not isinstance(values[COLUMN_ID], (int, float)):
                continue
            other_columns.update(header for header in edited if header not in APPLIED_COLUMNS)
            cells = {header: values[header] for header in edited if header in APPLIED_COLUMNS}
            if cells:
                # the sub category only helps to tell apart item categories with the same name
                cells[COLUMN_SUB_CATEGORY] = values.get(COLUMN_SUB_CATEGORY)
                edits[int(values[COLUMN_ID])] = cells
        workbook.close()
        return edits, other_columns

    def apply(self, product, cells):
        changes = {}

        if COLUMN_NAME in cells:
            name = self.corrected_name(product, cells)
            current = product["name_ru"] or product["name"]
            if not name:
                self.stats["names skipped (empty)"] += 1
                self.warn(f"product {product['id']}: empty name, skipped")
            elif name == clean(current):
                self.stats["names already applied"] += 1
            elif name in self.new_duplicates:
                self.stats["names skipped (another product would get the same name)"] += 1
                self.warn(f"product {product['id']}: {name!r} would repeat another product's name, skipped")
            else:
                if name != clean(cells[COLUMN_NAME]):
                    self.stats["names: code left out (it lives in articul_code)"] += 1
                # the columns that repeat the Russian name follow it, a real translation stays
                changes.update({
                    field: name for field in NAME_FIELDS if field == "name_ru" or product[field] == current
                })
                self.stats["names changed"] += 1
                self.log(product, "name", current, name)

        if COLUMN_BRAND in cells:
            brand = self.get_brand(clean(cells[COLUMN_BRAND]))
            if not brand:
                self.stats["brands skipped (empty)"] += 1
                self.warn(f"product {product['id']}: empty brand, skipped")
            elif brand.pk == product["brand_id"]:
                self.stats["brands already applied"] += 1
            else:
                changes.update(self.brand_change(product, brand))
                self.stats["brands changed"] += 1
                self.log(product, "brand", product["brand_id"], f"{brand.pk} {brand.name}")

        if COLUMN_ITEM_CATEGORY in cells:
            item_category = self.get_item_category(
                clean(cells[COLUMN_ITEM_CATEGORY]), clean(cells[COLUMN_SUB_CATEGORY]),
                product["product_item_category__product_sub_category__product_category_id"],
            )
            if not item_category:
                self.stats["item categories skipped (not found / ambiguous)"] += 1
                self.warn(f"product {product['id']}: item category {cells[COLUMN_ITEM_CATEGORY]!r} "
                          f"is not found or ambiguous, skipped")
            elif item_category.pk == product["product_item_category_id"]:
                self.stats["item categories already applied"] += 1
            else:
                changes.update(self.item_category_change(product, item_category))
                self.stats["item categories changed"] += 1
                self.log(product, "item category", product["product_item_category_id"],
                         f"{item_category.pk} {item_category.name}")

        if changes:
            Product.objects.rewrite(False).filter(pk=product["id"]).update(**changes)
            self.stats["products changed"] += 1

    def brand_change(self, product, brand):
        changes = {"brand": brand}
        # a product's model belongs to its brand (same rule as the admin API)
        if product["product_model__brand_id"] not in (None, brand.pk):
            changes["product_model"] = None
        self.touched_item_categories.add(product["product_item_category_id"])
        return changes

    def item_category_change(self, product, item_category):
        # what belongs to the old category goes: attribute values (as in the admin API) and range filter values
        ProductAttributeValue.objects.filter(product_id=product["id"]).exclude(
            attribute__category_links__item_category=item_category
        ).delete()
        ProductFilterNumericValue.objects.filter(product_id=product["id"]).exclude(
            schema__item_category=item_category
        ).delete()
        self.touched_item_categories.update((product["product_item_category_id"], item_category.pk))
        self.moved_to.add(item_category.pk)
        return {"product_item_category": item_category}

    def log(self, product, field, old, new):
        if self.verbose:
            self.stdout.write(f"  product {product['id']} {field}: {old!r} -> {new!r}")

    def get_brand(self, name):
        if not name:
            return None
        key = name.lower()
        if key not in self.brands:
            brand = named(Brand, name).first()
            if not brand:
                # no logo yet: hidden from the public brand list until someone uploads it in the admin
                brand = Brand.objects.create(name=name, name_ru=name, image="", is_visible=False)
                self.stats["brands created"] += 1
                self.stdout.write(f"  Brand {brand.pk} {name!r} created (hidden, without a logo)")
            self.brands[key] = brand
        return self.brands[key]

    def get_item_category(self, name, sub_category_name, category_id):
        key = (name.lower(), sub_category_name.lower(), category_id)
        if key not in self.item_categories:
            found = list(named(ProductItemCategory, name).select_related("product_sub_category")) if name else []
            if len(found) > 1 and sub_category_name:
                sub_category_ids = set(named(type(found[0].product_sub_category), sub_category_name).values_list(
                    "id", flat=True))
                found = [item for item in found if item.product_sub_category_id in sub_category_ids]
            if len(found) > 1:
                found = [item for item in found if item.product_sub_category.product_category_id == category_id]
            self.item_categories[key] = found[0] if len(found) == 1 else None
        return self.item_categories[key]
