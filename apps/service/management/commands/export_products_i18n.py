import html
import os
import re

import openpyxl
from django.conf import settings
from django.core.management.base import BaseCommand
from django.db.models import Exists, OuterRef
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from apps.service.models import Product, ProductImage

TAG_RE = re.compile(r"<[a-zA-Z/][^>]*>")
# Excel cell limit is 32767 characters
MAX_CELL = 32000

COLUMNS = [
    ("ID", "id", 8),
    ("Nomi (uz)", "name_uz", 40),
    ("Название (ru)", "name_ru", 40),
    ("Kategoriya (uz)", "cat_uz", 22),
    ("Категория (ru)", "cat_ru", 22),
    ("Subkategoriya (uz)", "sub_uz", 24),
    ("Подкатегория (ru)", "sub_ru", 24),
    ("Predmet kategoriya (uz)", "item_uz", 26),
    ("Предметная категория (ru)", "item_ru", 26),
    ("Qisqa tavsif (uz)", "short_description_uz", 45),
    ("Краткое описание (ru)", "short_description_ru", 45),
    ("Tavsif (uz)", "description_uz", 60),
    ("Описание (ru)", "description_ru", 60),
    ("Rasm bor", "has_image", 10),
]


def html_to_text(value):
    if not value:
        return ""
    if TAG_RE.search(value):
        value = re.sub(r"(?i)<br\s*/?>|</p>|</li>|</h\d>|</div>|</tr>", "\n", value)
        value = re.sub(r"(?i)<li[^>]*>", "• ", value)
        value = re.sub(r"<[^>]+>", "", value)
    value = html.unescape(value).replace("\xa0", " ")
    value = re.sub(r"[ \t]+\n", "\n", value)
    value = re.sub(r"\n{3,}", "\n\n", value).strip()
    return ILLEGAL_CHARACTERS_RE.sub("", value)[:MAX_CELL]


class Command(BaseCommand):
    help = "Export products to Excel: id, uz/ru name, category chain, short description, description, has image"

    def add_arguments(self, parser):
        parser.add_argument(
            "--output",
            type=str,
            default="products_i18n.xlsx",
            help="Output file path (default: products_i18n.xlsx in project root)",
        )
        parser.add_argument(
            "--raw-html",
            action="store_true",
            help="Keep descriptions as stored HTML instead of plain text",
        )

    def handle(self, *args, **options):
        output_path = options["output"]
        if not os.path.isabs(output_path):
            output_path = os.path.join(settings.BASE_DIR, output_path)
        clean = (lambda v: ILLEGAL_CHARACTERS_RE.sub("", v or "")[:MAX_CELL]) if options["raw_html"] else html_to_text

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Products"
        ws.append([title for title, _, _ in COLUMNS])
        for idx, (_, _, width) in enumerate(COLUMNS, start=1):
            ws.cell(1, idx).font = Font(bold=True)
            ws.column_dimensions[get_column_letter(idx)].width = width

        qs = (
            Product.objects.select_related("product_item_category__product_sub_category__product_category")
            .annotate(
                has_image=Exists(
                    ProductImage.objects.filter(product=OuterRef("pk")).exclude(image="")
                )
            )
            .order_by("id")
        )

        total = qs.count()
        without_image = 0
        self.stdout.write(f"Exporting {total} products...")
        wrap = Alignment(wrap_text=True, vertical="top")

        for product in qs.iterator(chunk_size=500):
            item_cat = product.product_item_category
            sub_cat = item_cat.product_sub_category if item_cat else None
            cat = sub_cat.product_category if sub_cat else None
            if not product.has_image:
                without_image += 1

            ws.append([
                product.id,
                clean(product.name_uz),
                clean(product.name_ru),
                (cat.name_uz or "") if cat else "",
                (cat.name_ru or "") if cat else "",
                (sub_cat.name_uz or "") if sub_cat else "",
                (sub_cat.name_ru or "") if sub_cat else "",
                (item_cat.name_uz or "") if item_cat else "",
                (item_cat.name_ru or "") if item_cat else "",
                clean(product.short_description_uz),
                clean(product.short_description_ru),
                clean(product.description_uz),
                clean(product.description_ru),
                "Ha" if product.has_image else "Yo'q",
            ])
            for cell in ws[ws.max_row]:
                cell.alignment = wrap

        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        wb.save(output_path)
        self.stdout.write(self.style.SUCCESS(f"\nDone! File saved to: {output_path}"))
        self.stdout.write(self.style.SUCCESS(f"Rows exported: {total} | Without image: {without_image}"))
