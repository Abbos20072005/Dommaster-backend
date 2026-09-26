from .models import ProductItemCategory, ProductSubCategory, ProductCategory, Product
import requests
import threading
import json
import logging
from html import escape
from io import BytesIO
from django.conf import settings
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)


def build_breadcrumbs(obj):
    breadcrumbs = []

    if isinstance(obj, Product):
        breadcrumbs.insert(0, {"id": obj.id, "name": obj.name, "level": "product"})
        obj = obj.product_item_category

    if isinstance(obj, ProductItemCategory):
        breadcrumbs.insert(0, {"id": obj.id, "name": obj.name, "level": "product_item_category"})
        obj = obj.product_sub_category

    if isinstance(obj, ProductSubCategory):
        breadcrumbs.insert(0, {"id": obj.id, "name": obj.name, "level": "product_sub_category"})
        obj = obj.product_category

    if isinstance(obj, ProductCategory):
        breadcrumbs.insert(0, {"id": obj.id, "name": obj.name, "level": "product_category"})

    return breadcrumbs


TELEGRAM_CAPTION_LIMIT = 1024


def _format_price(value):
    return f"{float(value or 0):,.0f}".replace(",", " ")


ORDER_STATUS_RU = {
    0: "В ожидании",
    1: "В процессе сборки",
    2: "В процессе доставки",
    3: "Выполнен",
    4: "Отменен",
    "0": "В ожидании",
    "1": "В процессе сборки",
    "2": "В процессе доставки",
    "3": "Выполнен",
    "4": "Отменен",
    "pending": "В ожидании",
    "collecting": "В процессе сборки",
    "delivering": "В процессе доставки",
    "completed": "Выполнен",
    "canceled": "Отменен",
    "cancelled": "Отменен",
}

PAYMENT_STATUS_RU = {
    0: "В ожидании",
    1: "Холд",
    2: "Оплачено",
    3: "Отменено",
    "0": "В ожидании",
    "1": "Холд",
    "2": "Оплачено",
    "3": "Отменено",
    "pending": "В ожидании",
    "hold": "Холд",
    "paid": "Оплачено",
    "cancelled": "Отменено",
    "canceled": "Отменено",
}

UNIT_RU = {
    "pcs": "шт.",
    "kg": "кг",
    "l": "л",
    "sm": "см",
    "m": "м",
    "g": "г",
    "pkg": "уп.",
    "set": "компл.",
}


def _get_russian_order_status(status_val):
    if status_val is None:
        return "В ожидании"
    if isinstance(status_val, int) and status_val in ORDER_STATUS_RU:
        return ORDER_STATUS_RU[status_val]
    s = str(status_val).lower().strip()
    return ORDER_STATUS_RU.get(s, str(status_val))


def _get_russian_payment_status(status_val):
    if status_val is None:
        return "В ожидании"
    if isinstance(status_val, int) and status_val in PAYMENT_STATUS_RU:
        return PAYMENT_STATUS_RU[status_val]
    s = str(status_val).lower().strip()
    return PAYMENT_STATUS_RU.get(s, str(status_val))


def _get_russian_unit(unit_val):
    if not unit_val:
        return ""
    s = str(unit_val).lower().strip()
    return UNIT_RU.get(s, str(unit_val))


def _build_order_caption(order_data):
    header = [
        f"<b>🛒 Новый заказ № {order_data['id']}</b>",
    ]

    # Customer info
    customer_role = order_data.get("customer_role")
    role_badge = " <i>[👷 Прораб]</i>" if customer_role and str(customer_role).lower() == "prorab" else ""
    header.append(f"👤 Клиент: {escape(order_data['customer_name'] or '—')}{role_badge}")
    header.append(f"📞 Телефон: {escape(order_data['customer_phone'] or '—')}")

    # Delivery info
    if order_data.get("delivery_type") == "Самовывоз":
        branch = escape(order_data.get("pickup_branch") or "—")
        header.append(f"🏬 Самовывоз: <b>{branch}</b>")
    else:
        addr = escape(order_data.get("delivery_address") or "—")
        header.append(f"🚚 Доставка: <b>{addr}</b>")
        if order_data.get("map_url"):
            header.append(f"🗺 Локация: <a href=\"{order_data['map_url']}\">Открыть на карте</a>")

    # Payment & Order Status info
    payment_type = order_data.get("payment_type") or "—"
    payment_method = f" ({order_data['payment_method']})" if order_data.get("payment_method") else ""
    payment_status = _get_russian_payment_status(order_data.get("payment_status"))
    order_status = _get_russian_order_status(order_data.get("order_status"))
    header.append(f"💳 Оплата: {escape(payment_type)}{escape(payment_method)} — <b>{escape(payment_status)}</b>")
    header.append(f"📋 Статус: <b>{escape(order_status)}</b>")

    # Total weight (if > 0)
    total_weight = order_data.get("total_weight", 0)
    if total_weight > 0:
        header.append(f"⚖️ Общий вес: <b>{total_weight:,.2f} кг</b>".replace(",", " "))

    # Promocode & Discount
    if order_data.get("promocode"):
        header.append(f"🏷 Промокод: <code>{escape(order_data['promocode'])}</code> (Скидка: {_format_price(order_data['saved_price'])} сум)")

    # Total Price
    header.append(f"💰 Итого к оплате: <b>{_format_price(order_data['total_price'])} сум</b>")
    header.append("\n📦 <b>Товары:</b>")

    caption = "\n".join(header)

    items = order_data["items"]
    for index, item in enumerate(items):
        unit_str = f" {item.get('unit')}" if item.get('unit') else ""
        line = f"\n • {escape(item['name'])} — {item['quantity']}{unit_str}"
        rest = len(items) - index - 1
        tail = f"\n… и ещё {rest} (см. Excel)" if rest else ""
        if len(caption) + len(line) + len(tail) > TELEGRAM_CAPTION_LIMIT:
            caption += f"\n… и ещё {len(items) - index} (см. Excel)"
            break
        caption += line

    return caption


def _style_range(ws, cell_range, font=None, fill=None, border=None, alignment=None):
    """Apply styles to all cells in a range or merged range."""
    for row in ws[cell_range]:
        for cell in row:
            if font is not None:
                cell.font = font
            if fill is not None:
                cell.fill = fill
            if border is not None:
                cell.border = border
            if alignment is not None:
                cell.alignment = alignment


def _build_order_excel(order_data):
    wb = Workbook()
    ws = wb.active
    ws.title = f"Заказ {order_data['id']}"
    ws.views.sheetView[0].showGridLines = True

    # Color Palette (Corporate Modern)
    NAVY_PRIMARY = "1E3A8A"      # Main Brand Navy
    NAVY_DARK = "0F2766"         # Subheader Dark Navy
    CARD_HEADER_BG = "F1F5F9"    # Light Gray for Card Labels
    CARD_VALUE_BG = "FFFFFF"     # Clean White for Values
    ZEBRA_BG = "F8FAFC"          # Table Alternating Row
    TOTAL_BG = "EFF6FF"          # Soft Blue for Grand Total
    BORDER_COLOR = "CBD5E1"      # Subtle Gray Border
    TEXT_DARK = "0F172A"         # Slate 900
    TEXT_MUTED = "475569"        # Slate 600
    TEXT_BLUE = "2563EB"         # Link Blue
    SUCCESS_GREEN = "16A34A"     # Green for Discount

    # Fonts
    font_title = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    font_subtitle = Font(name="Calibri", size=10, color="E2E8F0")
    font_card_lbl = Font(name="Calibri", size=10, bold=True, color=TEXT_MUTED)
    font_card_val = Font(name="Calibri", size=10, color=TEXT_DARK)
    font_card_val_bold = Font(name="Calibri", size=10, bold=True, color=TEXT_DARK)
    font_link = Font(name="Calibri", size=10, bold=True, color=TEXT_BLUE, underline="single")
    font_th = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_td = Font(name="Calibri", size=10, color=TEXT_DARK)
    font_td_bold = Font(name="Calibri", size=10, bold=True, color=TEXT_DARK)
    font_total_lbl = Font(name="Calibri", size=10, bold=True, color=TEXT_MUTED)
    font_total_val = Font(name="Calibri", size=11, bold=True, color=TEXT_DARK)
    font_grand_lbl = Font(name="Calibri", size=12, bold=True, color=NAVY_PRIMARY)
    font_grand_val = Font(name="Calibri", size=13, bold=True, color=NAVY_PRIMARY)
    font_discount = Font(name="Calibri", size=10, bold=True, color=SUCCESS_GREEN)

    # Fills
    fill_primary = PatternFill("solid", fgColor=NAVY_PRIMARY)
    fill_dark = PatternFill("solid", fgColor=NAVY_DARK)
    fill_card_lbl = PatternFill("solid", fgColor=CARD_HEADER_BG)
    fill_card_val = PatternFill("solid", fgColor=CARD_VALUE_BG)
    fill_zebra = PatternFill("solid", fgColor=ZEBRA_BG)
    fill_white = PatternFill("solid", fgColor="FFFFFF")
    fill_grand = PatternFill("solid", fgColor=TOTAL_BG)

    # Borders
    border_thin = Side(style="thin", color=BORDER_COLOR)
    border_cell = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
    border_double = Side(style="double", color=NAVY_PRIMARY)
    border_grand = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_double)

    # Set Column Widths first
    column_specs = [
        ("A", 6),   # №
        ("B", 15),  # Код 1С
        ("C", 14),  # Артикул
        ("D", 17),  # Штрихкод
        ("E", 15),  # Бренд
        ("F", 46),  # Наименование
        ("G", 8),   # Ед.
        ("H", 10),  # Кол-во
        ("I", 13),  # Вес ед., кг
        ("J", 14),  # Общ. вес, кг
        ("K", 16),  # Цена, сум
        ("L", 18),  # Сумма, сум
    ]
    for col_letter, width in column_specs:
        ws.column_dimensions[col_letter].width = width

    # 1. Header Banner (Rows 1 & 2)
    ws.row_dimensions[1].height = 32
    ws.merge_cells("A1:L1")
    ws["A1"] = f"   BUILDEX  |  ЗАКАЗ № {order_data['id']}"
    _style_range(ws, "A1:L1", font=font_title, fill=fill_primary, alignment=Alignment(vertical="center", horizontal="left"))

    order_status_ru = _get_russian_order_status(order_data.get("order_status"))
    payment_status_ru = _get_russian_payment_status(order_data.get("payment_status"))
    pay_display = order_data.get("payment_type") or "—"
    if order_data.get("payment_method"):
        pay_display += f" ({order_data['payment_method']})"

    ws.row_dimensions[2].height = 20
    ws.merge_cells("A2:L2")
    ws["A2"] = f"   Дата: {order_data['created_at'] or '—'}   |   Статус заказа: {order_status_ru}   |   Оплата: {pay_display} ({payment_status_ru})"
    _style_range(ws, "A2:L2", font=font_subtitle, fill=fill_dark, alignment=Alignment(vertical="center", horizontal="left"))

    ws.row_dimensions[3].height = 10

    # 2. Info Cards (Rows 4-8)
    role_str = " (Прораб)" if str(order_data.get("customer_role", "")).lower() == "prorab" else ""
    weight_str = f"{order_data.get('total_weight', 0):,.2f} кг".replace(",", " ") if order_data.get("total_weight") else "—"

    left_fields = [
        ("Клиент", f"{order_data['customer_name'] or '—'}{role_str}", True, None),
        ("Телефон", order_data["customer_phone"] or "—", False, None),
        ("Тип доставки", order_data["delivery_type"] or "—", False, None),
        ("Адрес", order_data["delivery_address"] or "—", False, None),
        ("Локация", "🗺 Открыть на карте (Google Maps)" if order_data.get("map_url") else "—", False, order_data.get("map_url")),
    ]

    right_fields = [
        ("Оплата", pay_display, False),
        ("Статус оплаты", payment_status_ru, True),
        ("Статус заказа", order_status_ru, True),
        ("Промокод", order_data["promocode"] or "—", False),
        ("Общий вес", weight_str, True),
    ]


    for idx, ((lbl_l, val_l, bold_l, link_l), (lbl_r, val_r, bold_r)) in enumerate(zip(left_fields, right_fields), start=4):
        ws.row_dimensions[idx].height = 22

        # Left Label (Merged A:B, width 21)
        range_lbl_l = f"A{idx}:B{idx}"
        ws.merge_cells(range_lbl_l)
        ws[f"A{idx}"] = lbl_l
        _style_range(ws, range_lbl_l, font=font_card_lbl, fill=fill_card_lbl, border=border_cell,
                     alignment=Alignment(vertical="center", horizontal="right"))

        # Left Value (Merged C:F, width 92)
        range_val_l = f"C{idx}:F{idx}"
        ws.merge_cells(range_val_l)
        ws[f"C{idx}"] = val_l
        if link_l:
            ws[f"C{idx}"].hyperlink = link_l
            _style_range(ws, range_val_l, font=font_link, fill=fill_card_val, border=border_cell,
                         alignment=Alignment(vertical="center", horizontal="left", indent=1))
        else:
            f_val = font_card_val_bold if bold_l else font_card_val
            _style_range(ws, range_val_l, font=f_val, fill=fill_card_val, border=border_cell,
                         alignment=Alignment(vertical="center", horizontal="left", indent=1))

        # Right Label (Merged G:H, width 18)
        range_lbl_r = f"G{idx}:H{idx}"
        ws.merge_cells(range_lbl_r)
        ws[f"G{idx}"] = lbl_r
        _style_range(ws, range_lbl_r, font=font_card_lbl, fill=fill_card_lbl, border=border_cell,
                     alignment=Alignment(vertical="center", horizontal="right"))

        # Right Value (Merged I:L, width 61)
        range_val_r = f"I{idx}:L{idx}"
        ws.merge_cells(range_val_r)
        ws[f"I{idx}"] = val_r
        f_val_r = font_card_val_bold if bold_r else font_card_val
        _style_range(ws, range_val_r, font=f_val_r, fill=fill_card_val, border=border_cell,
                     alignment=Alignment(vertical="center", horizontal="left", indent=1))

    ws.row_dimensions[9].height = 12

    # 3. Products Table Header (Row 10)
    headers = [
        "№", "Код 1С", "Артикул", "Штрихкод", "Бренд",
        "Наименование товара", "Ед.", "Кол-во",
        "Вес ед. (кг)", "Общ. вес (кг)", "Цена (сум)", "Сумма (сум)"
    ]
    ws.row_dimensions[10].height = 26
    for col_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(row=10, column=col_idx)
        cell.value = h_text
        cell.font = font_th
        cell.fill = fill_primary
        cell.border = border_cell
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)

    # 4. Products Rows (Rows 11+)
    curr_row = 11
    for item_idx, item in enumerate(order_data["items"], start=1):
        ws.row_dimensions[curr_row].height = 22
        row_fill = fill_zebra if item_idx % 2 == 1 else fill_white

        row_data = [
            (item_idx, "center", None, font_td),
            (item["product_code"], "center", None, font_td),
            (item["articul_code"], "center", None, font_td),
            (item.get("barcode", "—"), "center", None, font_td),
            (item.get("brand", "—"), "center", None, font_td),
            (item["name"], "left", None, font_td),
            (item["unit"], "center", None, font_td),
            (item["quantity"], "center", None, font_td_bold),
            (item.get("unit_weight") or 0.0, "right", "#,##0.00", font_td),
            (item.get("total_weight") or 0.0, "right", "#,##0.00", font_td),
            (item["price"], "right", "#,##0", font_td),
            (item["price"] * item["quantity"], "right", "#,##0", font_td_bold),
        ]

        for col_idx, (val, align, num_fmt, f) in enumerate(row_data, start=1):
            cell = ws.cell(row=curr_row, column=col_idx)
            cell.value = val
            cell.font = f
            cell.fill = row_fill
            cell.border = border_cell
            indent_val = 1 if align == "left" else 0
            cell.alignment = Alignment(vertical="center", horizontal=align, indent=indent_val)
            if num_fmt:
                cell.number_format = num_fmt

        curr_row += 1

    # Row space
    ws.row_dimensions[curr_row].height = 8
    curr_row += 1

    # 5. Totals Block
    totals = [
        ("Стоимость товаров:", order_data["products_total_price"], font_total_lbl, font_total_val, fill_white, False),
        ("Стоимость доставки:", order_data["delivery_price"], font_total_lbl, font_total_val, fill_white, False),
    ]

    if order_data.get("saved_price"):
        totals.append(("Скидка / Промокод:", -float(order_data["saved_price"]), font_discount, font_discount, fill_white, False))

    totals.append(("ИТОГО К ОПЛАТЕ:", order_data["total_price"], font_grand_lbl, font_grand_val, fill_grand, True))

    for label, val, f_lbl, f_val, r_fill, is_grand in totals:
        ws.row_dimensions[curr_row].height = 26 if is_grand else 21

        # Label (merged H:J)
        range_lbl = f"H{curr_row}:J{curr_row}"
        ws.merge_cells(range_lbl)
        ws[f"H{curr_row}"] = label
        b_style = border_grand if is_grand else border_cell
        _style_range(ws, range_lbl, font=f_lbl, fill=r_fill, border=b_style,
                     alignment=Alignment(vertical="center", horizontal="right"))

        # Value (merged K:L)
        range_val = f"K{curr_row}:L{curr_row}"
        ws.merge_cells(range_val)
        ws[f"K{curr_row}"] = float(val or 0)
        _style_range(ws, range_val, font=f_val, fill=r_fill, border=b_style,
                     alignment=Alignment(vertical="center", horizontal="right", indent=1))
        ws[f"K{curr_row}"].number_format = '#,##0" сум"'

        curr_row += 1

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()




def _send_telegram_message_sync(order_data):
    """Internal synchronous function that runs in a separate thread."""
    url = f"https://api.telegram.org/bot{settings.TELEGRAM_BOT_TOKEN}/sendDocument"
    payload = {
        "chat_id": settings.TELEGRAM_CHANNEL_ID,
        "caption": _build_order_caption(order_data),
        "parse_mode": "HTML",
    }
    if getattr(settings, "TELEGRAM_TOPIC_ID", None):
        payload["message_thread_id"] = settings.TELEGRAM_TOPIC_ID

    if order_data.get("map_url"):
        payload["reply_markup"] = json.dumps({
            "inline_keyboard": [
                [{"text": "🗺 Открыть на карте", "url": order_data["map_url"]}]
            ]
        })

    try:
        files = {
            "document": (
                f"order_{order_data['id']}.xlsx",
                _build_order_excel(order_data),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        }
        response = requests.post(url, data=payload, files=files, timeout=20)
        data = response.json()
        if not data.get("ok"):
            logger.error(
                "Telegram sendDocument failed: order_id=%s status=%s error_code=%s description=%s",
                order_data["id"], response.status_code, data.get("error_code"), data.get("description"),
            )
    except requests.RequestException as e:
        logger.error("Telegram request error: order_id=%s %s", order_data["id"], type(e).__name__)
    except Exception:
        logger.exception("Failed to send order to Telegram: order_id=%s", order_data["id"])


def send_telegram_message(order):
    """Send Telegram notification (caption + Excel) asynchronously to avoid blocking the request."""

    customer = order.customer
    customer_name = order.receiver_name or (customer.full_name if customer else None)
    customer_phone = order.receiver_phone or (customer.phone_number if customer else None)
    customer_role = customer.role if customer else None

    address = order.order_location
    delivery_address = None
    map_url = None
    if address:
        parts = [p for p in (address.name, address.location_name) if p]
        delivery_address = ", ".join(parts) if parts else None
        if address.latitude and address.longitude:
            map_url = f"https://maps.google.com/?q={address.latitude},{address.longitude}"

    order_items = order.order_items.select_related('product', 'product__brand').all()
    items_info = []
    total_order_weight = 0.0
    for item in order_items:
        prod = item.product
        unit_weight = float(prod.weight) if prod.weight else 0.0
        line_weight = round(unit_weight * item.quantity, 3)
        total_order_weight += line_weight
        items_info.append({
            'name': prod.name,
            'quantity': item.quantity,
            'product_code': prod.product_code or "—",
            'articul_code': prod.articul_code or "—",
            'barcode': prod.barcode or "—",
            'brand': prod.brand.name if prod.brand else "—",
            'unit': _get_russian_unit(prod.unit or prod.get_unit_display()),
            'unit_weight': unit_weight,
            'total_weight': line_weight,
            'price': prod.discount_price if prod.discount_price is not None else prod.price,
        })
    total_order_weight = round(total_order_weight, 3)

    order_data = {
        'id': order.id,
        'created_at': order.created_at.strftime("%d.%m.%Y %H:%M") if order.created_at else None,
        'customer_name': customer_name,
        'customer_phone': customer_phone,
        'customer_role': customer_role,
        'delivery_address': delivery_address,
        'map_url': map_url,
        'delivery_type': "Самовывоз" if order.delivery_type == 1 else "Доставка",
        'pickup_branch': order.pickup_branch.name if order.pickup_branch else None,
        'payment_type': order.get_payment_type_display() if order.payment_type else None,
        'payment_method': order.get_payment_method_display() if order.payment_method else None,
        'payment_status': _get_russian_payment_status(order.payment_status),
        'order_status': _get_russian_order_status(order.status),
        'promocode': order.promocode.code if order.promocode else None,
        'products_total_price': order.products_total_price,
        'delivery_price': float(order.delivery_price or 0),
        'saved_price': order.saved_price,
        'total_price': order.total_price,
        'total_weight': total_order_weight,
        'items': items_info,
    }

    thread = threading.Thread(
        target=_send_telegram_message_sync,
        args=(order_data,),
        daemon=True
    )
    thread.start()

