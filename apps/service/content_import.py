"""Bulk product content in an XLSX workbook: export (current content + the empty fields), import = upload →
background validation report (`ContentImport` + its `ContentImportChange` rows) → apply, all rows in one
transaction. Only the fields the admin panel owns are written; the 1C ones are in the file for reference and
ignored on import. An empty cell never changes anything (clear a field in the product form)."""
import logging
import re
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.conf import settings
from django.db import transaction
from django.db.models import Q
from django.utils import timezone
from django.utils.html import linebreaks
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.service.models import ContentImport, ContentImportChange, Product

logger = logging.getLogger(__name__)

SHEET = "Контент"
HELP_SHEET = "Инструкция"
# translated fields the admin panel owns (1C sends name / description only when it creates a product)
TEXT_FIELDS = ("name", "short_description", "description", "meta_title", "meta_description", "slug")
HTML_FIELDS = ("short_description", "description")
NUMBER_FIELDS = ("weight", "length", "width", "height")
LABELS = {
    "id": "ID продукта (ключ строки)",
    "product_code": "Код 1С",
    "articul_code": "Артикул",
    "barcode": "Штрихкод",
    "brand": "Бренд",
    "item_category": "Категория / подкатегория / предметная категория",
    "price": "Цена",
    "quantity": "Остаток",
    "is_active": "Активен",
    "publish_status": "Статус публикации",
    "missing": "Незаполненные поля",
    "name": "Название",
    "short_description": "Краткое описание (HTML или текст)",
    "description": "Описание (HTML или текст)",
    "meta_title": "Meta title",
    "meta_description": "Meta description",
    "slug": "Slug (латиница, цифры, дефис)",
    "weight": "Вес, кг",
    "length": "Длина, м",
    "width": "Ширина, м",
    "height": "Высота, м",
}
# owned by 1C (or computed): in the file for reference only
ONE_C_COLUMNS = ("product_code", "articul_code", "barcode", "brand", "item_category", "price", "quantity",
                 "is_active")
READ_ONLY_COLUMNS = ("id",) + ONE_C_COLUMNS + ("publish_status", "missing")
MAX_FILE_SIZE = 20 * 1024 * 1024
MAX_ROWS = 50_000
MAX_ERRORS = 500
SLUG_RE = re.compile(r"^[a-z0-9]+(?:-[a-z0-9]+)*$")
TAG_RE = re.compile(r"<[a-zA-Z/!]")
NUMBER_STEP = Decimal("0.001")
NUMBER_LIMIT = Decimal("10000000")  # DecimalField(max_digits=10, decimal_places=3)

READ_ONLY_FILL = PatternFill("solid", fgColor="D9D9D9")
HEADER_FILL = PatternFill("solid", fgColor="C6E0B4")
EMPTY_FILL = PatternFill("solid", fgColor="FFF2CC")


class ImportFileError(Exception):
    pass


def languages():
    return list(settings.MODELTRANSLATION_LANGUAGES)


def editable_columns(langs=None):
    """The fields an import may write, in the file's column order."""
    langs = langs or languages()
    return [f"{field}_{lang}" for field in TEXT_FIELDS for lang in langs] + list(NUMBER_FIELDS)


def column_label(column):
    field, _, lang = column.rpartition("_")
    if field in TEXT_FIELDS and lang in languages():
        return f"{LABELS[field]} ({lang})"
    return LABELS.get(column, column)


def as_text(column, value):
    """A stored value of an editable column as text ("" = empty)."""
    if column in NUMBER_FIELDS:
        return "" if value is None else str(Decimal(value).quantize(NUMBER_STEP))
    return value or ""


def current_value(product, column):
    return as_text(column, getattr(product, column))


def comparable(value):
    # what survives a trip through Excel: no control characters, \n line ends, no outer spaces
    return ILLEGAL_CHARACTERS_RE.sub("", value).replace("\r\n", "\n").replace("\r", "\n").strip()


def incomplete_q(columns):
    q = Q()
    for column in columns:
        q |= Q(**{f"{column}__isnull": True})
        if column not in NUMBER_FIELDS:
            q |= Q(**{column: ""})
    return q


# --- export ---

def _plain(value):
    # what Excel keeps: no control characters, \n line ends (openpyxl turns \r\n into two line breaks)
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value).replace("\r\n", "\n").replace("\r", "\n")
    return value


def _cell(ws, value, fill=None, bold=False):
    value = _plain(value)
    cell = WriteOnlyCell(ws, value=value)
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"  # text, not a formula
    if fill:
        cell.fill = fill
    if bold:
        cell.font = Font(bold=True)
    return cell


def export_workbook(products, langs=None, incomplete=False):
    """XLSX bytes: one row per product (`products` = a queryset), the read-only columns grey and marked
    `[1C]`, the empty editable cells yellow and listed in `missing`; the second sheet explains the columns."""
    columns = editable_columns(langs)
    if incomplete:
        products = products.filter(incomplete_q(columns))
    item = "product_item_category"
    path = (f"{item}__product_sub_category__product_category__name_ru", f"{item}__product_sub_category__name_ru",
            f"{item}__name_ru")
    # plain values: model instances made the export noticeably slower
    products = products.order_by("id").values(
        "id", "product_code", "articul_code", "barcode", "brand__name_ru", *path, "price", "quantity", "is_active",
        "publish_status", *columns)

    wb = Workbook(write_only=True)
    ws = wb.create_sheet(SHEET)
    ws.freeze_panes = "B2"
    for index, column in enumerate(READ_ONLY_COLUMNS + tuple(columns)):
        width = 12 if column in READ_ONLY_COLUMNS or column in NUMBER_FIELDS else 40
        ws.column_dimensions[get_column_letter(index + 1)].width = width
    ws.append([_cell(ws, f"{c} [1C]" if c in ONE_C_COLUMNS else c, READ_ONLY_FILL, True) for c in READ_ONLY_COLUMNS]
              + [_cell(ws, c, HEADER_FILL, True) for c in columns])

    empty_counts = dict.fromkeys(columns, 0)
    total = 0
    for product in products.iterator(chunk_size=2000):
        total += 1
        values = {column: as_text(column, product[column]) for column in columns}
        missing = [column for column, value in values.items() if value == ""]
        for column in missing:
            empty_counts[column] += 1
        read_only = [
            product["id"], product["product_code"] or "", product["articul_code"] or "", product["barcode"] or "",
            product["brand__name_ru"] or "", " / ".join(product[name] for name in path if product[name]),
            product["price"], product["quantity"], "да" if product["is_active"] else "нет",
            product["publish_status"], ", ".join(missing),
        ]
        row = read_only + [None if value == "" else Decimal(value) if column in NUMBER_FIELDS else value
                           for column, value in values.items()]
        ws.append([_cell(ws, value) if isinstance(value, str) and value.startswith("=") else _plain(value)
                   for value in row])

    # styles by rule, not per cell (per-cell styles made the export several times slower)
    last_row = total + 1
    read_only_range = f"A2:{get_column_letter(len(READ_ONLY_COLUMNS))}{last_row}"
    first = get_column_letter(len(READ_ONLY_COLUMNS) + 1)
    editable_range = f"{first}2:{get_column_letter(len(READ_ONLY_COLUMNS) + len(columns))}{last_row}"
    if total:
        ws.conditional_formatting.add(read_only_range, FormulaRule(formula=["TRUE"], fill=READ_ONLY_FILL))
        ws.conditional_formatting.add(editable_range,
                                      FormulaRule(formula=[f"LEN(TRIM({first}2))=0"], fill=EMPTY_FILL))

    help_ws = wb.create_sheet(HELP_SHEET)
    help_ws.column_dimensions["A"].width = 24
    help_ws.column_dimensions["B"].width = 50
    help_ws.column_dimensions["C"].width = 30
    help_ws.column_dimensions["D"].width = 16
    for line in (
        f"Продуктов в файле: {total}. Строка = продукт, ключ строки — колонка id (не менять).",
        "Меняются только зелёные колонки. Серые ([1C] и служебные) — только для чтения: импорт их не читает.",
        "Пустая ячейка = поле не меняется (жёлтые ячейки — незаполненные поля). Очистить поле — в карточке продукта.",
        "Описания: HTML или обычный текст (абзацы — переносами строк). Колонки языка можно удалить — "
        "импорт меняет только те колонки, что есть в файле.",
        "Импорт: файл загружается, проверяется, показывается отчёт; применяются сразу все строки или ни одной.",
        "",
    ):
        help_ws.append([_cell(help_ws, line)])
    help_ws.append([_cell(help_ws, title, HEADER_FILL, True)
                    for title in ("Колонка", "Что это", "Импорт", "Пустых")])
    for column in READ_ONLY_COLUMNS:
        kind = "1C, только чтение" if column in ONE_C_COLUMNS else "только чтение"
        help_ws.append([_cell(help_ws, column, READ_ONLY_FILL), _cell(help_ws, column_label(column)),
                        _cell(help_ws, kind)])
    for column in columns:
        help_ws.append([_cell(help_ws, column), _cell(help_ws, column_label(column)),
                        _cell(help_ws, "меняется"), _cell(help_ws, empty_counts[column])])

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --- validation ---

def _text(raw):
    if raw is None:
        return ""
    if isinstance(raw, float) and raw.is_integer():
        raw = int(raw)
    return str(raw).strip()


def clean_value(column, raw):
    """`raw` (a non-empty cell) → the value to store as text; ValueError with the message when it does not fit."""
    if column in NUMBER_FIELDS:
        if isinstance(raw, bool):
            raise ValueError("Ожидается число")
        try:
            number = Decimal(str(raw).strip().replace(",", ".").replace(" ", ""))
        except InvalidOperation:
            raise ValueError("Ожидается число")
        if not number.is_finite() or number < 0:
            raise ValueError("Ожидается неотрицательное число")
        number = number.quantize(NUMBER_STEP)
        if number >= NUMBER_LIMIT:
            raise ValueError("Слишком большое число")
        return str(number)
    value = _text(raw)
    field = column.rpartition("_")[0]
    if field == "slug":
        value = value.lower()
        if not SLUG_RE.match(value):
            raise ValueError("Slug: только латиница, цифры и дефис")
    max_length = Product._meta.get_field(column).max_length
    if max_length and len(value) > max_length:
        raise ValueError(f"Длиннее {max_length} символов")
    return value


def _header_key(value):
    return _text(value).split(" [")[0].strip().lower()


def read_sheet(path):
    """(editable columns {index: column}, ignored column titles, rows [(row number, id cell, {column: raw})])."""
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        raise ImportFileError("Файл не открывается как XLSX")
    try:
        ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None) or ()
        keys = [_header_key(value) for value in header]
        if "id" not in keys:
            raise ImportFileError("Нет колонки id")
        id_index = keys.index("id")
        allowed = set(editable_columns())
        columns = {index: key for index, key in enumerate(keys) if key in allowed}
        if not columns:
            raise ImportFileError("Нет ни одной редактируемой колонки")
        if len(set(columns.values())) != len(columns):
            raise ImportFileError("Колонка повторяется")
        ignored = [_text(value) for index, value in enumerate(header)
                   if _text(value) and index != id_index and index not in columns]
        data = []
        for number, row in enumerate(rows, start=2):
            if not any(_text(value) for value in row):
                continue
            if len(data) >= MAX_ROWS:
                raise ImportFileError(f"Больше {MAX_ROWS} строк")
            cells = {column: row[index] for index, column in columns.items() if index < len(row)}
            data.append((number, row[id_index] if id_index < len(row) else None, cells))
        return columns, ignored, data
    finally:
        wb.close()


def build_report(path):
    """Reads the file and checks every row against the database. Returns the report fields of `ContentImport`
    and the `ContentImportChange` rows (unsaved) of the rows without errors."""
    columns, ignored, data = read_sheet(path)
    errors = []
    parsed = []
    seen = {}
    for number, raw_id, cells in data:
        text_id = _text(raw_id)
        if not text_id.isdigit():
            errors.append({"row": number, "product": None, "column": "id", "message": "Нет id продукта"})
            continue
        product_id = int(text_id)
        if product_id in seen:
            errors.append({"row": number, "product": product_id, "column": "id",
                           "message": f"Продукт уже есть в строке {seen[product_id]}"})
            continue
        seen[product_id] = number
        parsed.append((number, product_id, cells))

    used = sorted(set(columns.values()), key=editable_columns().index)
    products = Product.objects.only("id", *used).in_bulk(list(seen))
    row_changes = {}
    row_errors = defaultdict(list)
    new_slugs = defaultdict(dict)  # column → {slug: (row, product id)}
    for number, product_id, cells in parsed:
        product = products.get(product_id)
        if product is None:
            row_errors[number].append({"row": number, "product": product_id, "column": "id",
                                       "message": "Продукт не найден"})
            continue
        changes = []
        for column, raw in cells.items():
            if _text(raw) == "":
                continue
            old = current_value(product, column)
            if column not in NUMBER_FIELDS and comparable(_text(raw)) == comparable(old):
                continue
            try:
                new = clean_value(column, raw)
            except ValueError as e:
                row_errors[number].append({"row": number, "product": product_id, "column": column,
                                           "message": str(e)})
                continue
            if column.rpartition("_")[0] in HTML_FIELDS and not TAG_RE.search(new):
                new = linebreaks(new)
            if comparable(new) == comparable(old):
                continue
            if column.startswith("slug_"):
                other = new_slugs[column].get(new)
                if other:
                    row_errors[number].append({"row": number, "product": product_id, "column": column,
                                               "message": f"Такой же slug в строке {other[0]}"})
                    continue
                new_slugs[column][new] = (number, product_id)
            changes.append(ContentImportChange(product_id=product_id, row=number, field=column,
                                               old_value=old, new_value=new))
        row_changes[number] = changes

    for column, slugs in new_slugs.items():
        taken = Product.objects.filter(**{f"{column}__in": list(slugs)}).values_list(column, "id")
        for slug, owner_id in taken:
            number, product_id = slugs[slug]
            if owner_id != product_id:
                row_errors[number].append({"row": number, "product": product_id, "column": column,
                                           "message": f"Slug занят продуктом {owner_id}"})

    for number, items in row_errors.items():
        errors.extend(items)
        row_changes.pop(number, None)
    errors.sort(key=lambda error: error["row"])
    error_rows = {error["row"] for error in errors}
    changes = [change for items in row_changes.values() for change in items]
    rows_changed = sum(1 for items in row_changes.values() if items)
    return {
        "columns": used,
        "languages": [lang for lang in languages() if any(c.endswith(f"_{lang}") for c in used)],
        "ignored_columns": ignored,
        "rows_total": len(data),
        "rows_changed": rows_changed,
        "rows_unchanged": len(row_changes) - rows_changed,
        "rows_error": len(error_rows),
        "changes_count": len(changes),
        "errors": errors[:MAX_ERRORS],
    }, changes


def validate_import(import_id):
    """Background step after the upload: the report + the change rows → status ready / invalid."""
    content_import = ContentImport.objects.filter(pk=import_id, status=ContentImport.VALIDATING).first()
    if not content_import:
        return
    try:
        report, changes = build_report(content_import.file.path)
    except ImportFileError as e:
        report, changes = {"error": str(e)}, []
    except Exception as e:
        logger.exception("Content import %s: validation failed", import_id)
        report, changes = {"error": f"Не удалось проверить файл: {e}"}, []
    has_errors = bool(report.get("error") or report.get("rows_error"))
    with transaction.atomic():
        content_import.changes.all().delete()
        if not has_errors:
            for change in changes:
                change.content_import = content_import
            ContentImportChange.objects.bulk_create(changes, batch_size=1000)
        for field, value in report.items():
            setattr(content_import, field, value)
        if has_errors:
            content_import.changes_count = 0
        content_import.status = ContentImport.INVALID if has_errors else ContentImport.READY
        content_import.save()


class ApplyConflict(Exception):
    def __init__(self, errors):
        super().__init__(f"{len(errors)} conflicts")
        self.errors = errors


def apply_import(import_id):
    """Writes every change of an import in one transaction: a product deleted or edited since the validation
    fails the whole import (nothing is written) — upload the file again."""
    content_import = ContentImport.objects.filter(pk=import_id, status=ContentImport.APPLYING).first()
    if not content_import:
        return
    try:
        with transaction.atomic():
            changes = list(content_import.changes.all())
            products = Product.objects.select_for_update().in_bulk({change.product_id for change in changes})
            fields = defaultdict(list)
            conflicts = []
            for change in changes:
                product = products.get(change.product_id)
                if product is None:
                    message = "Продукт удалён"
                elif comparable(current_value(product, change.field)) != comparable(change.old_value):
                    message = "Поле изменилось после проверки файла"
                else:
                    value = change.new_value
                    setattr(product, change.field, Decimal(value) if change.field in NUMBER_FIELDS else value)
                    fields[product].append(change.field)
                    continue
                conflicts.append({"row": change.row, "product": change.product_id, "column": change.field,
                                  "message": message})
            if conflicts:
                raise ApplyConflict(conflicts)
            for product, changed in fields.items():
                product.save(update_fields=changed + ["updated_at"])
            content_import.status = ContentImport.APPLIED
            content_import.applied_at = timezone.now()
            content_import.save(update_fields=["status", "applied_at", "updated_at"])
    except ApplyConflict as e:
        _fail(content_import, "Данные изменились после проверки — загрузите файл заново", e.errors)
    except Exception as e:
        logger.exception("Content import %s: apply failed", import_id)
        _fail(content_import, f"Не удалось применить: {e}", [])


def _fail(content_import, message, errors):
    content_import.status = ContentImport.FAILED
    content_import.error = message
    content_import.errors = errors[:MAX_ERRORS]
    content_import.save(update_fields=["status", "error", "errors", "updated_at"])
