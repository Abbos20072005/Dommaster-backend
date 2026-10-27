"""XLSX of the products' texts in both languages: id, uz / ru name, category chain, short description,
description (HTML -> plain text unless `raw_html`), has image. Read by `manage.py export_products_i18n` and the
admin `GET products/i18n-export/`."""
import html
import re
from io import BytesIO

from django.db.models import Exists, OuterRef
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from apps.service.models import ProductImage

TAG_RE = re.compile(r"<[a-zA-Z/][^>]*>")
# Excel cell limit is 32767 characters
MAX_CELL = 32000
SHEET = "Products"

COLUMNS = [
    ("ID", "id", 8),
    ("Nomi (uz)", "name_uz", 40),
    ("Название (ru)", "name_ru", 40),
    ("Kategoriya (uz)", "cat_uz", 22),
    ("Категория (ru)", "cat_ru", 22),
    ("Subkategoriya (uz)", "sub_uz", 24),
    ("Подкатегория (ru)", "sub_ru", 24),
    ("Predmet kategoriya (uz)", "item_uz", 26),
    ("Предметная категория (ru)", "item_ru", 26),
    ("Qisqa tavsif (uz)", "short_description_uz", 45),
    ("Краткое описание (ru)", "short_description_ru", 45),
    ("Tavsif (uz)", "description_uz", 60),
    ("Описание (ru)", "description_ru", 60),
    ("Rasm bor", "has_image", 10),
]

ITEM = "product_item_category"
SUB = f"{ITEM}__product_sub_category"
CAT = f"{SUB}__product_category"
CATEGORY_FIELDS = {
    "cat_uz": f"{CAT}__name_uz", "cat_ru": f"{CAT}__name_ru",
    "sub_uz": f"{SUB}__name_uz", "sub_ru": f"{SUB}__name_ru",
    "item_uz": f"{ITEM}__name_uz", "item_ru": f"{ITEM}__name_ru",
}
TEXT_FIELDS = ("name_uz", "name_ru", "short_description_uz", "short_description_ru", "description_uz",
               "description_ru")


def html_to_text(value):
    if not value:
        return ""
    if TAG_RE.search(value):
        value = re.sub(r"(?i)<br\s*/?>|</p>|</li>|</h\d>|</div>|</tr>", "\n", value)
        value = re.sub(r"(?i)<li[^>]*>", "• ", value)
        value = re.sub(r"<[^>]+>", "", value)
    value = html.unescape(value).replace("\xa0", " ")
    value = re.sub(r"[ \t]+\n", "\n", value)
    value = re.sub(r"\n{3,}", "\n\n", value).strip()
    return ILLEGAL_CHARACTERS_RE.sub("", value)[:MAX_CELL]


def raw_text(value):
    return ILLEGAL_CHARACTERS_RE.sub("", value or "")[:MAX_CELL]


def export_workbook(products, raw_html=False):
    """(XLSX bytes, rows count, rows without image) for `products` (a queryset), ordered by id."""
    clean = raw_text if raw_html else html_to_text
    rows = (
        products.annotate(has_image=Exists(ProductImage.objects.filter(product=OuterRef("pk")).exclude(image="")))
        .order_by("id")
        # plain values: model instances made the export noticeably slower
        .values("id", "has_image", *TEXT_FIELDS, *CATEGORY_FIELDS.values())
    )

    wb = Workbook(write_only=True)
    ws = wb.create_sheet(SHEET)
    ws.freeze_panes = "B2"
    for index, (_, _, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    bold = Font(bold=True)
    header = []
    for title, _, _ in COLUMNS:
        cell = WriteOnlyCell(ws, value=title)
        cell.font = bold
        header.append(cell)
    ws.append(header)

    wrap = Alignment(wrap_text=True, vertical="top")
    total = without_image = 0
    for product in rows.iterator(chunk_size=2000):
        total += 1
        if not product["has_image"]:
            without_image += 1
        values = {key: clean(product[key]) for key in TEXT_FIELDS}
        values.update({key: product[field] or "" for key, field in CATEGORY_FIELDS.items()})
        values["id"] = product["id"]
        values["has_image"] = "Ha" if product["has_image"] else "Yo'q"
        row = []
        for _, key, _ in COLUMNS:
            cell = WriteOnlyCell(ws, value=values[key])
            cell.alignment = wrap
            row.append(cell)
        ws.append(row)

    if total:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{total + 1}"
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue(), total, without_image
