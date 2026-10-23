"""Bulk content import / export (`apps/service/content_import.py`, admin `content-imports/`)."""
import shutil
import tempfile
from decimal import Decimal
from io import BytesIO
from unittest import mock

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.service import content_import
from apps.service.models import ContentImport, Product
from apps.service.tests import make_item_category, make_product

MEDIA = tempfile.mkdtemp()


def run_now(task, *args):
    task(*args)


def sheet_rows(data):
    ws = load_workbook(BytesIO(data))[content_import.SHEET]
    rows = list(ws.iter_rows(values_only=True))
    return list(rows[0]), rows[1:]


def edited(data, edits):
    """The exported workbook with `edits` = {(product id, column): value} written into it."""
    wb = load_workbook(BytesIO(data))
    ws = wb[content_import.SHEET]
    header = [cell.value for cell in ws[1]]
    for row in ws.iter_rows(min_row=2):
        for (product_id, column), value in edits.items():
            if row[0].value == product_id:
                row[header.index(column)].value = value
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@override_settings(MEDIA_ROOT=MEDIA)
@mock.patch("apps.service.api.v1.admin.views.content_import.run_in_background", run_now)
@mock.patch("apps.service.api.v1.admin.serializers.content_import.run_in_background", run_now)
class ContentImportTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(MEDIA, ignore_errors=True)

    def setUp(self):
        self.admin = get_user_model().objects.create_user("admin", password="x", is_staff=True)
        self.client = APIClient()
        self.client.force_authenticate(self.admin)
        _, item = make_item_category("Дрели")
        self.drill = make_product("D1", item, 100)
        self.drill.description_ru = "Строка 1\r\nСтрока 2"
        self.drill.meta_title_ru = "Дрель"
        self.drill.weight = Decimal("1.5")
        self.drill.save()
        self.saw = make_product("S1", item, 200)

    def export(self, **params):
        response = self.client.get("/api/v1/admin/products/content-export/", params)
        self.assertEqual(response.status_code, 200)
        return response.content

    def upload(self, data):
        with self.captureOnCommitCallbacks(execute=True):
            response = self.client.post("/api/v1/admin/content-imports/", {
                "file": SimpleUploadedFile("content.xlsx", data)}, format="multipart")
        self.assertEqual(response.status_code, 201, response.data)
        return ContentImport.objects.get(pk=response.data["id"])

    def apply(self, content_import_row):
        with self.captureOnCommitCallbacks(execute=True):
            response = self.client.post(f"/api/v1/admin/content-imports/{content_import_row.pk}/apply/")
        content_import_row.refresh_from_db()
        return response

    def test_export_marks_1c_columns_and_empty_fields(self):
        header, rows = sheet_rows(self.export())
        self.assertIn("price [1C]", header)
        self.assertIn("name_uz", header)
        drill = next(row for row in rows if row[0] == self.drill.pk)
        missing = drill[header.index("missing")]
        self.assertIn("meta_title_uz", missing)
        self.assertNotIn("meta_title_ru", missing)
        self.assertEqual(drill[header.index("weight")], 1.5)

    def test_export_one_language_and_incomplete_only(self):
        header, _ = sheet_rows(self.export(languages="uz"))
        self.assertIn("name_uz", header)
        self.assertNotIn("name_ru", header)
        self.assertEqual(self.client.get("/api/v1/admin/products/content-export/", {"languages": "en"}).status_code,
                         400)
        Product.objects.update(short_description_uz="x", short_description_ru="x", description_uz="x",
                               description_ru="x", meta_title_uz="x", meta_title_ru="x", meta_description_uz="x",
                               meta_description_ru="x", weight=1, length=1, width=1, height=1)
        _, rows = sheet_rows(self.export(incomplete="true"))
        self.assertEqual(rows, [])

    def test_unchanged_export_changes_nothing(self):
        row = self.upload(self.export())
        self.assertEqual(row.status, ContentImport.READY)
        self.assertEqual((row.rows_total, row.rows_unchanged, row.changes_count), (2, 2, 0))
        self.assertEqual(self.apply(row).status_code, 400)

    def test_import_writes_only_admin_fields_in_one_batch(self):
        data = edited(self.export(), {
            (self.drill.pk, "name_uz"): "Drel",
            (self.drill.pk, "description_uz"): "Birinchi\nIkkinchi",
            (self.drill.pk, "weight"): "2,25",
            (self.drill.pk, "price [1C]"): 1,
            (self.saw.pk, "slug_ru"): "pila-new",
        })
        row = self.upload(data)
        self.assertEqual(row.status, ContentImport.READY, row.errors)
        self.assertEqual((row.rows_changed, row.rows_unchanged, row.changes_count), (2, 0, 4))
        self.assertIn("price [1C]", row.ignored_columns)
        self.assertEqual(self.apply(row).status_code, 202)
        self.assertEqual(row.status, ContentImport.APPLIED)
        self.assertEqual(row.applied_by, self.admin)
        self.drill.refresh_from_db()
        self.saw.refresh_from_db()
        self.assertEqual(self.drill.name_uz, "Drel")
        self.assertEqual(self.drill.description_uz, "<p>Birinchi<br>Ikkinchi</p>")
        self.assertEqual(self.drill.weight, Decimal("2.250"))
        self.assertEqual(self.drill.price, 100)
        self.assertEqual(self.saw.slug_ru, "pila-new")
        changes = self.client.get(f"/api/v1/admin/content-imports/{row.pk}/changes/",
                                  {"product": self.drill.pk}).data
        self.assertEqual(changes["count"], 3)

    def test_any_error_blocks_the_whole_file(self):
        data = edited(self.export(), {
            (self.drill.pk, "name_uz"): "Drel",
            (self.saw.pk, "weight"): "abc",
            (self.saw.pk, "slug_uz"): self.drill.slug_uz,
        })
        row = self.upload(data)
        self.assertEqual(row.status, ContentImport.INVALID)
        self.assertEqual((row.rows_error, row.changes_count), (1, 0))
        self.assertEqual({error["column"] for error in row.errors}, {"weight", "slug_uz"})
        self.assertEqual(self.apply(row).status_code, 400)
        self.drill.refresh_from_db()
        self.assertNotEqual(self.drill.name_uz, "Drel")

    def test_unknown_and_repeated_ids(self):
        wb = load_workbook(BytesIO(self.export()))
        ws = wb[content_import.SHEET]
        header = [cell.value for cell in ws[1]]
        blank = [None] * len(header)
        for product_id in (999999, self.drill.pk):
            line = list(blank)
            line[0] = product_id
            line[header.index("name_uz")] = "X"
            ws.append(line)
        buffer = BytesIO()
        wb.save(buffer)
        row = self.upload(buffer.getvalue())
        self.assertEqual(row.status, ContentImport.INVALID)
        self.assertEqual([error["message"] for error in row.errors][-2:],
                         ["Продукт не найден", f"Продукт уже есть в строке 2"])

    def test_not_a_workbook(self):
        row = self.upload(b"not a zip")
        self.assertEqual(row.status, ContentImport.INVALID)
        self.assertTrue(row.error)

    def test_conflict_rolls_back_everything(self):
        row = self.upload(edited(self.export(), {(self.drill.pk, "name_uz"): "Drel",
                                                 (self.saw.pk, "name_uz"): "Arra"}))
        Product.objects.filter(pk=self.saw.pk).update(name_uz="Edited in the panel")
        self.apply(row)
        self.assertEqual(row.status, ContentImport.FAILED)
        self.assertEqual(row.errors[0]["product"], self.saw.pk)
        self.drill.refresh_from_db()
        self.assertNotEqual(self.drill.name_uz, "Drel")

    def test_applied_import_stays_in_the_log(self):
        row = self.upload(edited(self.export(), {(self.drill.pk, "name_uz"): "Drel"}))
        self.apply(row)
        self.assertEqual(self.client.delete(f"/api/v1/admin/content-imports/{row.pk}/").status_code, 400)
        listed = self.client.get("/api/v1/admin/content-imports/", {"status": "applied"}).data
        self.assertEqual(listed["count"], 1)
